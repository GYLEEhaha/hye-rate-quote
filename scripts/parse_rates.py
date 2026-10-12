#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AQ国际 运价 Excel 解析脚本
用法: python3 scripts/parse_rates.py
读取 data/rates.xlsx，输出 data/parsed/*.json
"""

from __future__ import annotations

import json
import re
from datetime import datetime, date
from pathlib import Path
from typing import Any, Optional

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "rates.xlsx"
OUT = ROOT / "data" / "parsed"

SKIP_VALUES = {"/", "\\", "-", "—", "无匹配项", "#N/A", "#REF!", "#VALUE!", None, ""}


def cell(ws: Worksheet, r: int, c: int) -> Any:
    return ws.cell(r, c).value


def s(v: Any) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    t = str(v).strip()
    if t in SKIP_VALUES or t.upper() in {"#N/A", "#REF!", "#VALUE!"}:
        return None
    return t


def num(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (v != v):  # NaN
            return None
        return float(v)
    t = str(v).strip().replace(",", "")
    if t in SKIP_VALUES or t.upper().startswith("#"):
        return None
    try:
        return float(t)
    except ValueError:
        return None


def is_price(v: Any) -> bool:
    return num(v) is not None


def looks_like_section_title(text: Optional[str]) -> bool:
    if not text:
        return False
    # zone / data row labels — never treat as product titles
    zone_labels = (
        "美国西岸", "美国中部", "美国东岸", "美国岛屿",
        "西岸", "中部", "东岸",
    )
    if text.strip() in zone_labels or text.strip() in ("1区", "2区", "3区", "4区", "5区", "6区"):
        return False
    keys = (
        "海派", "海卡", "商私卡", "超大件", "小包", "岛屿", "集运",
        "空派", "空卡", "头程",
        "极速达", "限时达", "准时达", "快提", "定提", "普船", "特惠",
        "美东", "美西", "美中", "HYE",
    )
    skip = (
        "比重", "特别提示", "附加", "交货仓", "体积", "预报", "地址类",
        "包装类", "签收", "扣件", "滞留", "拒收", "美国新政", "保险",
        "赔偿", "内置电池", "品名", "区域", "邮编", "仓库代码", "分类",
        "日截", "一截", "四截", "生效日期", "返回",
    )
    if any(k in text for k in skip):
        return False
    return any(k in text for k in keys) and len(text) < 80


def parse_weight_bracket(label: str) -> dict:
    """Parse labels like '12-50KG（包税）', '51+KG（包税）', '300KG+(自税)', '1CBM+（包税）', '1-5kg（一口价）'"""
    info: dict[str, Any] = {
        "label": label,
        "tax_mode": None,
        "unit": "元/KG",
        "weight_min_kg": None,
        "weight_max_kg": None,
        "cbm_min": None,
        "flat_rate": False,
    }
    if "包税" in label:
        info["tax_mode"] = "包税"
    elif "自税" in label:
        info["tax_mode"] = "自税"
    if "一口价" in label:
        info["flat_rate"] = True
        info["unit"] = "元/票"

    # CBM
    m = re.search(r"(\d+(?:\.\d+)?)\s*CBM\+?", label, re.I)
    if m:
        info["cbm_min"] = float(m.group(1))
        info["unit"] = "元/CBM"
        return info

    # weight range 12-50KG
    m = re.search(r"(\d+(?:\.\d+)?)\s*[-–~到至]\s*(\d+(?:\.\d+)?)\s*[Kk][Gg]", label)
    if m:
        info["weight_min_kg"] = float(m.group(1))
        info["weight_max_kg"] = float(m.group(2))
        return info

    # 51+KG / 300KG+ / 41kg+
    m = re.search(r"(\d+(?:\.\d+)?)\s*[Kk][Gg]\s*\+", label)
    if m:
        info["weight_min_kg"] = float(m.group(1))
        return info
    m = re.search(r"(\d+(?:\.\d+)?)\s*\+\s*[Kk][Gg]", label)
    if m:
        info["weight_min_kg"] = float(m.group(1))
        return info

    # 1-5kg
    m = re.search(r"(\d+(?:\.\d+)?)\s*[-–~]\s*(\d+(?:\.\d+)?)\s*[Kk][Gg]", label)
    if m:
        info["weight_min_kg"] = float(m.group(1))
        info["weight_max_kg"] = float(m.group(2))
        return info

    return info


def zone_from_zip_range(zip_range: Optional[str], zone_hint: Optional[str] = None) -> Optional[str]:
    if zone_hint:
        z = zone_hint.replace("美国", "").replace("岸", "").strip()
        mapping = {
            "西岸": "美西", "西": "美西", "美西": "美西",
            "中部": "美中", "中": "美中", "美中": "美中",
            "东岸": "美东", "东": "美东", "美东": "美东",
            "头程": "头程", "岛屿": "岛屿",
        }
        for k, v in mapping.items():
            if k in zone_hint:
                return v
    if not zip_range:
        return None
    zr = zip_range.replace(" ", "")
    if "96700" in zr or "岛屿" in zr:
        return "岛屿"
    if "80000" in zr or "80000-99999" in zr:
        return "美西"
    if "40000" in zr:
        return "美中"
    if "00000" in zr or "0-39999" in zr:
        return "美东"
    return None


def zip_to_zone(zip_code: str) -> Optional[str]:
    try:
        z = int(str(zip_code).strip()[:5])
    except ValueError:
        return None
    if 96700 <= z <= 99999 and z >= 96700:
        # Hawaii etc often 967xx-968xx; Alaska 995xx-999xx — island sheet uses 96700-99999
        pass
    if 0 <= z <= 39999:
        return "美东"
    if 40000 <= z <= 79999:
        return "美中"
    if 80000 <= z <= 99999:
        return "美西"
    return None


def make_rate(**kwargs) -> dict:
    base = {
        "category": None,
        "product": None,
        "channel_name": None,
        "warehouse": None,
        "tax_mode": None,
        "zone": None,
        "zip_range": None,
        "warehouse_codes": [],
        "weight_min_kg": None,
        "weight_max_kg": None,
        "cbm_min": None,
        "unit": "元/KG",
        "unit_price": None,
        "weight_bracket": None,
        "transit_time": None,
        "notes": None,
        "order_channels": None,
        "source_sheet": None,
        "flat_rate": False,
    }
    base.update(kwargs)
    return base


def save_json(name: str, data: Any) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / name
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"  wrote {path} ({len(data) if isinstance(data, list) else 'obj'})")


# ---------- parsers ----------

def parse_catalog(wb) -> dict:
    ws = wb["目录"]
    channels = []
    warehouses = []
    # Branding for this site is always AQ国际 (Excel may still say 环洋)
    company = "AQ国际"
    for r in range(3, 40):
        # Sea books put channel in D; air books put product in E
        name = s(cell(ws, r, 5)) or s(cell(ws, r, 4)) or s(cell(ws, r, 3))
        updated = s(cell(ws, r, 8)) or s(cell(ws, r, 7))
        if not name:
            continue
        if "进入" in name or name in ("美国海运", "美国空运", "仓   库   地   址", "离境口岸代码", "附    录"):
            continue
        if name in ("大包渠道", "小包渠道", "大件渠道", "空卡渠道", "空派自税渠道", "服务渠道"):
            continue
        if "仓：" in name or "仓:" in name:
            warehouses.append(name)
        elif "联系" in name or "电话" in name:
            continue
        elif updated or "HYE" in name or "海" in name or "空" in name or "集运" in name or "系统" in name:
            channels.append({"name": name, "updated": updated})
    wh_details = []
    for r in range(14, 32):
        t = s(cell(ws, r, 3))
        if t and ("仓" in t or "联系" in t):
            wh_details.append(t)
    return {
        "company": company,
        "brand": "AQ",
        "channels_index": channels,
        "warehouse_addresses": wh_details,
        "parsed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }



def parse_haitou(ws: Worksheet) -> list[dict]:
    """海派头程报价汇总 — headhaul only, 包税, Shanghai/Yiwu vs Bantan"""
    rates = []
    category = None
    channel = None
    for r in range(6, 15):
        cat = s(cell(ws, r, 2))
        ch = s(cell(ws, r, 3))
        zone = s(cell(ws, r, 4))
        if cat and not cat.startswith("比重") and not cat.startswith("特别") and "附加" not in cat:
            if cat in ("快船", "普船"):
                category = cat
            elif looks_like_section_title(cat):
                continue
        if ch:
            channel = ch
        if not channel:
            continue
        if zone and zone not in ("头程", "/"):
            # skip non-headhaul if any
            pass
        sh = num(cell(ws, r, 6))
        bt = num(cell(ws, r, 7))
        transit = s(cell(ws, r, 8))
        notes = s(cell(ws, r, 9))
        order = s(cell(ws, r, 10))
        for wh, price in (("上海仓", sh), ("义乌仓", sh), ("坂田仓", bt)):
            if price is None:
                continue
            # 上海/义乌 share same column; skip duplicate 义乌 when same — keep both for filter UX
            rates.append(make_rate(
                category="海派头程",
                product=channel,
                channel_name=channel,
                warehouse=wh,
                tax_mode="包税",
                zone="头程",
                zip_range=None,
                unit="元/KG",
                unit_price=price,
                weight_bracket="头程计费",
                transit_time=transit,
                notes=notes,
                order_channels=order,
                source_sheet="海派头程报价汇总",
                ship_type=category,
            ))
    return rates


def parse_jiyun(ws: Worksheet) -> list[dict]:
    """集运价格汇总 — weight tiers by warehouse"""
    rates = []
    # headers row5: cols 6-8 上海, 9-11 义乌, 12-14 坂田
    brackets = []
    for c, wh in [(6, "上海仓"), (7, "上海仓"), (8, "上海仓"),
                  (9, "义乌仓"), (10, "义乌仓"), (11, "义乌仓"),
                  (12, "坂田仓"), (13, "坂田仓"), (14, "坂田仓")]:
        lab = s(cell(ws, 5, c))
        if lab:
            brackets.append((c, wh, parse_weight_bracket(lab)))

    category = None
    channel = None
    for r in range(6, 35):
        cat = s(cell(ws, r, 2))
        ch = s(cell(ws, r, 3))
        zone_raw = s(cell(ws, r, 4))
        zip_r = s(cell(ws, r, 5))
        if cat and cat.startswith("比重"):
            continue
        if cat and cat.startswith("特别"):
            break
        if cat and cat not in ("特惠",) and looks_like_section_title(cat) is False:
            if cat in ("极速达", "限时达", "准时达", "快提", "普船", "特惠", "美东"):
                category = cat
        if cat == "特惠":
            category = "特惠"
        if ch:
            channel = ch
        if not zone_raw and not zip_r:
            continue
        if not channel:
            continue
        zone = zone_from_zip_range(zip_r, zone_raw)
        transit = s(cell(ws, r, 15))
        notes = s(cell(ws, r, 16))
        order = s(cell(ws, r, 17))
        for c, wh, br in brackets:
            price = num(cell(ws, r, c))
            if price is None:
                continue
            rates.append(make_rate(
                category="集运",
                product=channel,
                channel_name=channel,
                warehouse=wh,
                tax_mode=br.get("tax_mode") or "包税",
                zone=zone or zone_raw,
                zip_range=zip_r if zip_r and zip_r != "/" else None,
                unit=br["unit"],
                unit_price=price,
                weight_min_kg=br["weight_min_kg"],
                weight_max_kg=br["weight_max_kg"],
                cbm_min=br["cbm_min"],
                weight_bracket=br["label"],
                transit_time=transit,
                notes=notes,
                order_channels=order,
                source_sheet="集运价格汇总",
                ship_type=category,
            ))
    return rates


def parse_haipai_channels(ws: Worksheet) -> list[dict]:
    """HYE海派渠道汇总 — multi product blocks with zone+weight+tax"""
    rates = []
    product = None
    transit = None
    notes = None
    order = None
    brackets = []  # (col, warehouse, bracket_info)

    for r in range(1, min(ws.max_row, 120) + 1):
        c2 = s(cell(ws, r, 2))
        # new product section
        if c2 and looks_like_section_title(c2) and "区域" not in c2 and "比重" not in c2:
            if any(k in c2 for k in ("海派", "HYE", "极速达", "限时达", "准时达", "快提", "定提", "普船", "特惠", "美东")):
                product = c2
                transit = notes = order = None
                brackets = []
                continue
        if c2 and c2.startswith("特别提示"):
            break
        if c2 and c2.startswith("比重"):
            continue
        # header row with 上海仓/义乌仓
        if c2 == "区域":
            # next row has brackets
            brackets = []
            # detect warehouse groups from row r
            # cols 4-6 上海/义乌, 7-9 坂田 typically
            continue
        # bracket row
        lab4 = s(cell(ws, r, 4))
        if lab4 and ("KG" in lab4.upper() or "CBM" in lab4.upper() or "包税" in lab4 or "自税" in lab4):
            brackets = []
            for c, wh in [(4, "上海仓"), (5, "上海仓"), (6, "上海仓"),
                          (7, "坂田仓"), (8, "坂田仓"), (9, "坂田仓")]:
                lab = s(cell(ws, r, c))
                if lab:
                    brackets.append((c, wh, parse_weight_bracket(lab)))
            # also duplicate 上海 as 义乌
            extra = []
            for c, wh, br in brackets:
                if wh == "上海仓":
                    extra.append((c, "义乌仓", br))
            brackets = brackets + extra
            continue

        zone_raw = c2
        zip_r = s(cell(ws, r, 3))
        if not zone_raw or not zip_r:
            continue
        if "美国" not in zone_raw and zone_raw not in ("美国西岸", "美国中部", "美国东岸"):
            if not re.match(r"^(美国|美西|美中|美东)", zone_raw or ""):
                continue
        zone = zone_from_zip_range(zip_r, zone_raw)
        t = s(cell(ws, r, 10))
        n = s(cell(ws, r, 11))
        o = s(cell(ws, r, 12))
        if t:
            transit = t
        if n:
            notes = n
        if o:
            order = o
        if not product or not brackets:
            continue
        for c, wh, br in brackets:
            price = num(cell(ws, r, c))
            if price is None:
                continue
            rates.append(make_rate(
                category="海派",
                product=product,
                channel_name=product,
                warehouse=wh,
                tax_mode=br.get("tax_mode"),
                zone=zone,
                zip_range=zip_r,
                unit=br["unit"],
                unit_price=price,
                weight_min_kg=br["weight_min_kg"],
                weight_max_kg=br["weight_max_kg"],
                cbm_min=br["cbm_min"],
                weight_bracket=br["label"],
                transit_time=transit,
                notes=notes,
                order_channels=order,
                source_sheet="HYE海派渠道汇总",
            ))
    return rates


def parse_haika_sheet(ws: Worksheet, sheet_name: str, default_product: str) -> list[dict]:
    """Parse HYE快船海卡 / HYE普船海卡 / HYE美西特惠普船LA海卡 style sheets"""
    rates = []
    product = default_product
    transit = None
    notes = None
    order = None
    brackets = []

    for r in range(1, min(ws.max_row, 200) + 1):
        c2 = s(cell(ws, r, 2))
        if c2 and looks_like_section_title(c2) and "仓库" not in c2 and "比重" not in c2:
            if "海卡" in c2 or "HYE" in c2 or "限时达" in c2 or "准时达" in c2 or "快提" in c2 or "定提" in c2 or "普船" in c2 or "特惠" in c2:
                product = c2
                brackets = []
                continue
        if c2 and (c2.startswith("比重") or c2.startswith("特别提示") or c2.startswith("附加")):
            if c2.startswith("特别") or c2.startswith("附加"):
                # keep going for more sections; only break on 特别 at end
                if c2.startswith("特别提示"):
                    break
            continue
        if c2 == "仓库代码":
            continue
        # bracket labels
        lab4 = s(cell(ws, r, 4))
        if lab4 and ("KG" in lab4.upper() or "CBM" in lab4.upper() or "包税" in lab4 or "自税" in lab4):
            brackets = []
            # detect columns dynamically
            for c in range(4, 16):
                lab = s(cell(ws, r, c))
                if not lab:
                    continue
                # warehouse from row above
                wh_header = s(cell(ws, r - 1, c))
                # walk left for merged warehouse header
                if not wh_header:
                    for cc in range(c, 3, -1):
                        wh_header = s(cell(ws, r - 1, cc))
                        if wh_header:
                            break
                wh = "上海仓"
                if wh_header:
                    if "坂田" in wh_header:
                        wh = "坂田仓"
                    elif "义乌" in wh_header and "上海" not in wh_header:
                        wh = "义乌仓"
                    elif "上海" in wh_header:
                        wh = "上海仓"
                br = parse_weight_bracket(lab)
                brackets.append((c, wh, br))
                if wh == "上海仓" and "义乌" in (wh_header or "上海仓/义乌仓"):
                    brackets.append((c, "义乌仓", br))
            continue

        if not c2:
            continue
        # warehouse code row
        codes_raw = c2
        if codes_raw.startswith("比重") or codes_raw.startswith("特别"):
            continue
        codes = [x.strip() for x in re.split(r"[/／,，\s]+", codes_raw) if x.strip()]
        # filter non-code noise
        if not codes:
            continue
        # skip if looks like prose
        if len(codes_raw) > 80 and "海外仓" not in codes_raw:
            continue

        t = s(cell(ws, r, 12))
        claim = s(cell(ws, r, 13))
        n = s(cell(ws, r, 14))
        o = s(cell(ws, r, 15))
        if t:
            transit = t
        if claim and transit:
            transit = f"{transit}；{claim}"
        elif claim and not t:
            pass
        if n:
            notes = n
        if o:
            order = o

        if not brackets:
            continue

        for c, wh, br in brackets:
            price = num(cell(ws, r, c))
            if price is None:
                continue
            rates.append(make_rate(
                category="海卡",
                product=product,
                channel_name=product,
                warehouse=wh,
                tax_mode=br.get("tax_mode"),
                zone=None,
                zip_range=None,
                warehouse_codes=codes if "海外仓" not in codes_raw else ["海外仓自提/商业地址头程"],
                unit=br["unit"],
                unit_price=price,
                weight_min_kg=br["weight_min_kg"],
                weight_max_kg=br["weight_max_kg"],
                cbm_min=br["cbm_min"],
                weight_bracket=br["label"],
                transit_time=transit,
                notes=notes,
                order_channels=order,
                source_sheet=sheet_name,
            ))
    return rates


def parse_price_summary(ws: Worksheet) -> list[dict]:
    """价格汇总 — FBA warehouse exploded rows (best for warehouse code lookup)"""
    rates = []
    # row1 headers warehouses, row2 brackets
    # cols: 6-9 上海, 11-14 坂田, 15-18 义乌
    bracket_defs = [
        (6, "上海仓"), (7, "上海仓"), (8, "上海仓"), (9, "上海仓"),
        (11, "坂田仓"), (12, "坂田仓"), (13, "坂田仓"), (14, "坂田仓"),
        (15, "义乌仓"), (16, "义乌仓"), (17, "义乌仓"), (18, "义乌仓"),
    ]
    brackets = []
    for c, wh in bracket_defs:
        lab = s(cell(ws, 2, c))
        if lab:
            brackets.append((c, wh, parse_weight_bracket(lab)))

    for r in range(3, ws.max_row + 1):
        product_line = s(cell(ws, r, 1))
        channel = s(cell(ws, r, 2))
        wh_code = s(cell(ws, r, 3))
        wh_group = s(cell(ws, r, 4))
        transit = s(cell(ws, r, 19))
        if not wh_code or not channel:
            continue
        if wh_code.upper() in ("#N/A",):
            continue
        for c, wh, br in brackets:
            price = num(cell(ws, r, c))
            if price is None:
                continue
            rates.append(make_rate(
                category="海卡",
                product=product_line or channel,
                channel_name=channel,
                warehouse=wh,
                tax_mode=br.get("tax_mode"),
                zone=None,
                warehouse_codes=[wh_code],
                warehouse_group=wh_group,
                unit=br["unit"],
                unit_price=price,
                weight_min_kg=br["weight_min_kg"],
                weight_max_kg=br["weight_max_kg"],
                cbm_min=br["cbm_min"],
                weight_bracket=br["label"],
                transit_time=transit,
                source_sheet="价格汇总",
            ))
    return rates


def parse_shangsika(ws: Worksheet, sheet_name: str, category: str = "商私卡") -> list[dict]:
    """HYE美西商私卡 / HYE美中美东商私卡"""
    rates = []
    product = None
    transit = None
    notes = None
    order = None
    brackets = []
    mode = "zone"  # or fba

    for r in range(1, min(ws.max_row, 200) + 1):
        c2 = s(cell(ws, r, 2))
        if c2 and looks_like_section_title(c2) and "邮编" not in c2 and "比重" not in c2 and "非FBA" not in c2:
            if any(k in c2 for k in ("商私卡", "HYE", "限时达", "准时达", "快提", "定提", "普船", "美中", "美东", "休斯顿", "芝加哥")):
                product = c2
                brackets = []
                mode = "zone"
                continue
        if c2 and c2.startswith("比重"):
            continue
        if c2 and c2.startswith("特别提示"):
            break
        if c2 == "非FBA代码":
            mode = "fba"
            # re-read brackets from this row
            brackets = []
            for c in range(4, 18):
                lab = s(cell(ws, r, c))
                if not lab:
                    continue
                wh_header = None
                for cc in range(c, 3, -1):
                    wh_header = s(cell(ws, r - 1, cc)) if r > 1 else None
                    # actually headers are on earlier rows; use column ranges
                    break
                # approximate: 4-10 上海/义乌, 11-17 坂田
                wh = "上海仓" if c < 11 else "坂田仓"
                br = parse_weight_bracket(lab)
                brackets.append((c, wh, br))
                if wh == "上海仓":
                    brackets.append((c, "义乌仓", br))
            continue

        lab4 = s(cell(ws, r, 4))
        if lab4 and ("KG" in lab4.upper() or "CBM" in lab4.upper() or "包税" in lab4 or "自税" in lab4):
            if c2 in ("邮编分区", None) or s(cell(ws, r - 1, 2)) == "邮编分区" or True:
                brackets = []
                for c in range(4, 18):
                    lab = s(cell(ws, r, c))
                    if not lab:
                        continue
                    wh = "上海仓" if c < 11 else "坂田仓"
                    br = parse_weight_bracket(lab)
                    brackets.append((c, wh, br))
                    if wh == "上海仓":
                        brackets.append((c, "义乌仓", br))
                continue

        if not c2:
            continue
        if c2 in ("邮编分区", "/"):
            # zip prefix row for 美中
            zip_hint = s(cell(ws, r, 3))
            if zip_hint and brackets:
                t = s(cell(ws, r, 18))
                n = s(cell(ws, r, 19))
                o = s(cell(ws, r, 21))
                if t:
                    transit = t
                if n:
                    notes = n
                if o:
                    order = o
                for c, wh, br in brackets:
                    price = num(cell(ws, r, c))
                    if price is None:
                        continue
                    rates.append(make_rate(
                        category=category,
                        product=product or sheet_name,
                        channel_name=product,
                        warehouse=wh,
                        tax_mode=br.get("tax_mode"),
                        zone="美中" if "美中" in (product or "") else ("美东" if "美东" in (product or "") else "美西"),
                        zip_range=zip_hint,
                        unit=br["unit"],
                        unit_price=price,
                        weight_min_kg=br["weight_min_kg"],
                        weight_max_kg=br["weight_max_kg"],
                        cbm_min=br["cbm_min"],
                        weight_bracket=br["label"],
                        transit_time=transit,
                        notes=notes,
                        order_channels=order,
                        source_sheet=sheet_name,
                    ))
            continue

        # zone row like 1区 / FBA code
        zone_or_code = c2
        zip_hint = s(cell(ws, r, 3))
        t = s(cell(ws, r, 18))
        n = s(cell(ws, r, 19))
        o = s(cell(ws, r, 21))
        if t:
            transit = t
        if n:
            notes = n
        if o:
            order = o

        if not product or not brackets:
            continue

        is_fba = mode == "fba" or (
            zone_or_code and not zone_or_code.endswith("区") and "一件代发" not in zone_or_code
            and zone_or_code not in ("1区", "2区", "3区", "4区", "5区", "6区")
            and not zone_or_code.startswith("美西一件")
        )

        for c, wh, br in brackets:
            price = num(cell(ws, r, c))
            if price is None:
                continue
            rec = make_rate(
                category=category,
                product=product,
                channel_name=product,
                warehouse=wh,
                tax_mode=br.get("tax_mode"),
                zone="美西" if "美西" in sheet_name or "美西" in (product or "") else (
                    "美中" if "美中" in (product or "") or "休斯顿" in (product or "") or "芝加哥" in (product or "") else (
                        "美东" if "美东" in (product or "") else None
                    )
                ),
                zip_range=zip_hint,
                unit=br["unit"],
                unit_price=price,
                weight_min_kg=br["weight_min_kg"],
                weight_max_kg=br["weight_max_kg"],
                cbm_min=br["cbm_min"],
                weight_bracket=br["label"],
                transit_time=transit,
                notes=notes,
                order_channels=order,
                source_sheet=sheet_name,
                region_label=zone_or_code,
            )
            if is_fba and zone_or_code and not zone_or_code.endswith("区") and "一件代发" not in zone_or_code:
                rec["warehouse_codes"] = [zone_or_code]
            rates.append(rec)
    return rates


def parse_shangsika_zishui(ws: Worksheet) -> list[dict]:
    """HYE美西商私卡自税汇总 — CBM self-tax"""
    rates = []
    product = None
    transit = None
    notes = None
    order = None
    brackets = []

    for r in range(1, min(ws.max_row, 120) + 1):
        c1 = s(cell(ws, r, 1))
        if c1 and looks_like_section_title(c1) and "邮编" not in c1 and "比重" not in c1:
            product = c1
            brackets = []
            continue
        if c1 and c1.startswith("比重"):
            continue
        if c1 and c1.startswith("特别"):
            break
        lab3 = s(cell(ws, r, 3))
        if lab3 and ("CBM" in lab3.upper() or "自税" in lab3):
            brackets = []
            for c, wh in [(3, "上海仓"), (5, "上海仓"), (9, "坂田仓"), (11, "坂田仓")]:
                lab = s(cell(ws, r, c))
                if lab:
                    br = parse_weight_bracket(lab)
                    br["tax_mode"] = "自税"
                    brackets.append((c, wh, br))
                    if wh == "上海仓":
                        brackets.append((c, "义乌仓", br))
            continue
        if not c1 or c1 == "邮编分区":
            continue
        zip_hint = s(cell(ws, r, 2))
        t = s(cell(ws, r, 15))
        n = s(cell(ws, r, 16))
        o = s(cell(ws, r, 18))
        if t:
            transit = t
        if n:
            notes = n
        if o:
            order = o
        if not product or not brackets:
            continue
        for c, wh, br in brackets:
            price = num(cell(ws, r, c))
            if price is None:
                continue
            rates.append(make_rate(
                category="商私卡",
                product=product,
                channel_name=product,
                warehouse=wh,
                tax_mode="自税",
                zone="美西",
                zip_range=zip_hint,
                unit=br["unit"],
                unit_price=price,
                cbm_min=br["cbm_min"],
                weight_bracket=br["label"],
                transit_time=transit,
                notes=notes,
                order_channels=order,
                source_sheet="HYE美西商私卡自税汇总",
                region_label=c1,
            ))
    return rates


def parse_zone_simple(ws: Worksheet, sheet_name: str, category: str) -> list[dict]:
    """超大件 / 小包裹 / 岛屿海派 — simple zone x warehouse"""
    rates = []
    product = None
    transit = None
    notes = None
    order = None
    brackets = []  # (col, wh, br)
    zip_col = 3
    price_cols = [(4, "上海仓"), (5, "坂田仓")]

    # detect layout from header
    for r in range(1, 10):
        if s(cell(ws, r, 2)) == "区域":
            if s(cell(ws, r, 3)) and "邮编" in (s(cell(ws, r, 3)) or ""):
                zip_col = 3
                price_cols = [(4, "上海仓"), (5, "坂田仓")]
            elif s(cell(ws, r, 4)) and "邮编" in (s(cell(ws, r, 4)) or ""):
                zip_col = 4
                price_cols = [(6, "上海仓"), (7, "坂田仓")]
            break

    for r in range(1, min(ws.max_row, 80) + 1):
        c2 = s(cell(ws, r, 2))
        if not c2:
            # still may be bracket-only row
            pass

        if c2 and c2.startswith("特别提示"):
            break
        if c2 and (c2.startswith("比重") or c2.startswith("附加")):
            continue
        if c2 == "区域":
            continue

        # product title
        if c2 and looks_like_section_title(c2):
            product = c2
            brackets = []
            transit = notes = order = None
            continue

        # bracket / weight label row
        br_found = False
        for c, _wh in price_cols:
            lab = s(cell(ws, r, c))
            if lab and (("kg" in lab.lower()) or ("一口价" in lab) or ("KG" in lab)):
                brackets = []
                for c2_, wh in price_cols:
                    lab2 = s(cell(ws, r, c2_))
                    if not lab2:
                        continue
                    br = parse_weight_bracket(lab2)
                    if category == "小包裹":
                        br["flat_rate"] = True
                        br["unit"] = "元/票"
                        br["tax_mode"] = br.get("tax_mode")  # usually none
                    brackets.append((c2_, wh, br))
                    if wh == "上海仓":
                        brackets.append((c2_, "义乌仓", br))
                br_found = True
                break
        if br_found:
            continue

        if not c2:
            continue
        if not any(k in c2 for k in ("美国", "西岸", "中部", "东岸", "岛屿")):
            continue

        zip_r = s(cell(ws, r, zip_col))
        zone = zone_from_zip_range(zip_r, c2)

        # transit / notes / order — scan row
        for cc in range(6, 13):
            val = s(cell(ws, r, cc))
            if not val:
                continue
            if ("天" in val or "时效" in val or "开船" in val) and not transit:
                transit = val
            elif "HYE" in val:
                order = val
            elif len(val) > 8:
                notes = val

        if not product or not brackets:
            continue

        for c, wh, br in brackets:
            price = num(cell(ws, r, c))
            if price is None:
                continue
            rates.append(make_rate(
                category=category,
                product=product,
                channel_name=product,
                warehouse=wh,
                tax_mode=br.get("tax_mode") or ("包税" if category != "小包裹" else None),
                zone=zone,
                zip_range=zip_r if zip_r and re.search(r"\d{5}", zip_r or "") else None,
                unit=br["unit"],
                unit_price=price,
                weight_min_kg=br["weight_min_kg"],
                weight_max_kg=br["weight_max_kg"],
                weight_bracket=br["label"],
                transit_time=transit,
                notes=notes,
                order_channels=order,
                source_sheet=sheet_name,
                flat_rate=br.get("flat_rate", False),
            ))
    return rates



def _zone_row_ok(zone_raw: Optional[str]) -> bool:
    if not zone_raw:
        return False
    if zone_raw.startswith("比重") or zone_raw.startswith("特别") or "附加" in zone_raw:
        return False
    return True


def parse_kongpai(ws: Worksheet, sheet_name: str = "HYE空派", category: str = "空派", tax_mode: str = "包税") -> list[dict]:
    """HYE空派 / similar: product sections with 西/中/东/岛屿 × weight brackets × 上海/坂田."""
    rates: list[dict] = []
    product = None
    transit = notes = order = None
    # Detect bracket header row dynamically
    max_r = min(ws.max_row or 80, 120)
    r = 1
    while r <= max_r:
        b = s(cell(ws, r, 2))
        if b and looks_like_section_title(b):
            product = b
            transit = notes = order = None
            r += 1
            continue
        if b and (b.startswith("特别提示") or b.startswith("附加费用")):
            break
        # Header row with 区域 / 机场
        if b in ("区域", "机场") or (b and "区域" in b and s(cell(ws, r, 3)) in ("邮编", None)):
            # Next row may be weight brackets under warehouse groups
            # Find warehouse group starts from this header row
            wh_spans: list[tuple[str, int, int]] = []
            # Scan header row for warehouse names
            for c in range(4, 20):
                lab = s(cell(ws, r, c))
                if lab and ("仓" in lab or lab in ("上海仓", "坂田仓", "义乌仓", "嘉兴仓")):
                    wh_spans.append((lab if lab.endswith("仓") else lab, c, c))
            # Expand spans using next row brackets
            br_row = r + 1
            brackets_by_wh: list[tuple[str, int, dict]] = []
            if wh_spans:
                # Determine column ranges between warehouse starts
                starts = [c for _, c, _ in wh_spans] + [20]
                for i, (wh, start, _) in enumerate(wh_spans):
                    end = starts[i + 1]
                    # If next row has weight labels under this span, use them; else single col
                    found = False
                    for c in range(start, end):
                        lab = s(cell(ws, br_row, c))
                        if lab and ("KG" in lab.upper() or "kg" in lab or "CBM" in lab.upper() or "一口价" in lab):
                            brackets_by_wh.append((wh, c, parse_weight_bracket(lab)))
                            found = True
                    if not found:
                        # single price column (e.g. 自税 101+KG on header sub-row)
                        lab = s(cell(ws, br_row, start)) or "标准"
                        brackets_by_wh.append((wh, start, parse_weight_bracket(lab)))
            else:
                # Fallback: D=上海 E=坂田 style
                for c, wh in ((4, "上海仓"), (5, "坂田仓")):
                    lab = s(cell(ws, br_row, c)) or s(cell(ws, r, c)) or "标准"
                    if "KG" in (lab or "").upper() or "kg" in (lab or "") or lab == "标准" or "一口价" in (lab or ""):
                        brackets_by_wh.append((wh, c, parse_weight_bracket(lab)))

            # Data rows after bracket row
            rr = br_row + 1
            while rr <= max_r:
                zone_raw = s(cell(ws, rr, 2))
                if zone_raw and looks_like_section_title(zone_raw):
                    break
                if zone_raw and (zone_raw.startswith("比重") or zone_raw.startswith("特别") or zone_raw.startswith("附加")):
                    break
                if zone_raw and zone_raw in ("区域", "机场"):
                    break
                zip_r = s(cell(ws, rr, 3))
                # airport rows: B=芝加哥, C=60191
                if zone_raw and _zone_row_ok(zone_raw):
                    zone = zone_from_zip_range(zip_r if zip_r and "-" in str(zip_r) else None, zone_raw)
                    if not zone and zip_r and zip_r.isdigit():
                        zone = "头程"
                        # keep zip as airport code hint
                    # pull transit/notes/order from first data row of section if present
                    # find rightmost note columns
                    for c in range(20, 3, -1):
                        v = s(cell(ws, rr, c))
                        if not v:
                            continue
                        # heuristic: long text = transit or notes
                        if "工作日" in v and not transit:
                            transit = v
                        elif ("HYE" in v or "下单" in v) and not order:
                            order = v
                        elif len(v) > 8 and not notes:
                            notes = v
                    # Better: use known layouts
                    # 空派: J=transit K=notes L=order  (cols 10,11,12)
                    t2 = s(cell(ws, rr, 10)) or s(cell(ws, rr, 6)) or s(cell(ws, rr, 12))
                    n2 = s(cell(ws, rr, 11)) or s(cell(ws, rr, 7)) or s(cell(ws, rr, 13))
                    o2 = s(cell(ws, rr, 12)) or s(cell(ws, rr, 8)) or s(cell(ws, rr, 14))
                    if t2 and ("工作日" in t2 or "提取" in t2 or "签收" in t2):
                        transit = t2
                    if n2 and "HYE-" not in (n2 or "")[:4]:
                        notes = n2
                    if o2 and ("HYE" in o2 or "\\n" in o2):
                        order = o2
                    # Also check col 6/7/8 for 自税/超大件 layout
                    for cand in (s(cell(ws, rr, 6)), s(cell(ws, rr, 7)), s(cell(ws, rr, 8))):
                        if not cand:
                            continue
                        if "工作日" in cand:
                            transit = cand
                        elif cand.startswith("HYE") or "HYE-" in cand:
                            order = cand
                        elif len(cand) > 10 and "元" not in cand[:3]:
                            notes = notes or cand

                    for wh, c, br in brackets_by_wh:
                        price = num(cell(ws, rr, c))
                        if price is None:
                            continue
                        tm = br.get("tax_mode") or tax_mode
                        rates.append(make_rate(
                            category=category,
                            product=product or sheet_name,
                            channel_name=product or sheet_name,
                            warehouse=wh,
                            tax_mode=tm,
                            zone=zone or zone_raw,
                            zip_range=zip_r if zip_r and zip_r != "/" else None,
                            unit=br.get("unit") or "元/KG",
                            unit_price=price,
                            weight_bracket=br.get("label"),
                            weight_min_kg=br.get("weight_min_kg"),
                            weight_max_kg=br.get("weight_max_kg"),
                            cbm_min=br.get("cbm_min"),
                            transit_time=transit,
                            notes=notes,
                            order_channels=order,
                            source_sheet=sheet_name,
                            flat_rate=br.get("flat_rate", False),
                        ))
                rr += 1
            r = rr
            continue
        r += 1
    return rates


def parse_kongpai_xiaobao(ws: Worksheet, sheet_name: str = "HYE空派小包") -> list[dict]:
    """Flat-rate small parcel by billed weight tiers."""
    rates: list[dict] = []
    product = None
    transit = notes = order = None
    max_r = min(ws.max_row or 100, 150)
    r = 1
    while r <= max_r:
        b = s(cell(ws, r, 2))
        if b and looks_like_section_title(b):
            product = b
            transit = notes = order = None
            r += 1
            continue
        if b and (b.startswith("特别提示") or b.startswith("附加费用")):
            break
        if b == "区域" or (b and "区域" in b):
            # skip header + 一口价 subheader
            r += 1
            if s(cell(ws, r, 4)) == "一口价" or (s(cell(ws, r, 4)) and "一口价" in (s(cell(ws, r, 4)) or "")):
                r += 1
            prev_w = 0.0
            while r <= max_r:
                zone_raw = s(cell(ws, r, 2))
                if zone_raw and looks_like_section_title(zone_raw):
                    break
                if zone_raw and (zone_raw.startswith("特别") or zone_raw.startswith("附加") or zone_raw.startswith("比重")):
                    break
                if zone_raw == "区域":
                    break
                w = num(cell(ws, r, 3))
                sh = num(cell(ws, r, 4))
                bt = num(cell(ws, r, 5))
                if zone_raw and ("美国" in zone_raw or zone_raw == "美国"):
                    t2 = s(cell(ws, r, 6))
                    n2 = s(cell(ws, r, 7))
                    o2 = s(cell(ws, r, 9)) or s(cell(ws, r, 8))
                    if t2:
                        transit = t2
                    if n2:
                        notes = n2
                    if o2:
                        order = o2
                if w is None or (sh is None and bt is None):
                    r += 1
                    continue
                br_label = f"{w:g}kg（一口价）"
                for wh, price in (("上海仓", sh), ("坂田仓", bt)):
                    if price is None:
                        continue
                    rates.append(make_rate(
                        category="空派小包",
                        product=product or sheet_name,
                        channel_name=product or sheet_name,
                        warehouse=wh,
                        tax_mode="包税",
                        zone="岛屿" if (product and "岛屿" in product) else None,
                        zip_range=None,
                        unit="元/票",
                        unit_price=price,
                        weight_bracket=br_label,
                        weight_min_kg=prev_w,
                        weight_max_kg=w,
                        transit_time=transit,
                        notes=notes,
                        order_channels=order,
                        source_sheet=sheet_name,
                        flat_rate=True,
                    ))
                prev_w = w
                r += 1
            continue
        r += 1
    return rates


def parse_warehouse_surcharges(ws: Worksheet) -> list[dict]:
    """美国海外仓快速中转服务 — treat as surcharge list."""
    items = []
    for r in range(3, min(ws.max_row or 40, 50) + 1):
        name = s(cell(ws, r, 1))
        fee = s(cell(ws, r, 4)) or s(cell(ws, r, 6))
        note = s(cell(ws, r, 8))
        if name:
            items.append({"name": name, "fee": fee, "description": note})
    return items



def parse_schedule(ws: Worksheet) -> list[dict]:
    rows = []
    for r in range(5, 21):
        # East China
        name = s(cell(ws, r, 1))
        if name and name not in ("注意事项", "仓库地址"):
            rows.append({
                "region": "华东",
                "channel": name,
                "pol": s(cell(ws, r, 2)),
                "etd": s(cell(ws, r, 3)),
                "pod": s(cell(ws, r, 4)),
                "eta": s(cell(ws, r, 5)),
                "carrier": s(cell(ws, r, 6)),
                "vessel": s(cell(ws, r, 7)),
                "cutoff_shanghai": s(cell(ws, r, 8)),
                "cutoff_yiwu": s(cell(ws, r, 9)),
            })
        name2 = s(cell(ws, r, 15))
        if name2 and name2 not in ("注意事项", "仓库地址"):
            rows.append({
                "region": "华南",
                "channel": name2,
                "pol": s(cell(ws, r, 16)),
                "etd": s(cell(ws, r, 17)),
                "pod": s(cell(ws, r, 18)),
                "eta": s(cell(ws, r, 19)),
                "carrier": s(cell(ws, r, 20)),
                "vessel": s(cell(ws, r, 21)),
                "cutoff_bantan": s(cell(ws, r, 22)),
            })
    return rows


def parse_surcharges(ws: Worksheet) -> list[dict]:
    items = []
    for r in range(2, min(ws.max_row, 50) + 1):
        name = s(cell(ws, r, 1))
        fee = s(cell(ws, r, 2))
        desc = s(cell(ws, r, 3))
        if name:
            items.append({"name": name, "fee": fee, "description": desc})
    return items


def assign_ids(rates: list[dict]) -> list[dict]:
    for i, r in enumerate(rates):
        r["id"] = f"R{i+1:05d}"
    return rates


def build_meta(catalog: dict, rates: list[dict], schedule: list, surcharges: list) -> dict:
    cats = sorted({r["category"] for r in rates if r.get("category")})
    products_by_cat: dict[str, list] = {}
    for r in rates:
        c = r.get("category") or "其他"
        p = r.get("product") or r.get("channel_name")
        if not p:
            continue
        products_by_cat.setdefault(c, [])
        if p not in products_by_cat[c]:
            products_by_cat[c].append(p)
    warehouses = sorted({r["warehouse"] for r in rates if r.get("warehouse")})
    tax_modes = sorted({r["tax_mode"] for r in rates if r.get("tax_mode")})
    # unique FBA codes
    fba = sorted({
        code
        for r in rates
        for code in (r.get("warehouse_codes") or [])
        if code and "海外仓" not in code and len(code) <= 20
    })
    return {
        **catalog,
        "categories": cats,
        "products_by_category": products_by_cat,
        "warehouses": warehouses,
        "tax_modes": tax_modes,
        "fba_warehouse_codes": fba,
        "rate_count": len(rates),
        "schedule": schedule,
        "surcharges": surcharges,
        "zip_zones": {
            "美东": "00000-39999",
            "美中": "40000-79999",
            "美西": "80000-99999",
            "岛屿": "96700-99999（岛屿渠道）",
        },
        "disclaimer": "本系统价格仅供内部参考，不作为最终结算依据。实际价格以对应渠道价格表及客服确认为准。",
    }


def main() -> None:
    print(f"Loading {XLSX} ...")
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    catalog = parse_catalog(wb)
    all_rates: list[dict] = []

    # Sea book (legacy) + air book (current). Missing sheets are skipped.
    parsers = [
        ("海派头程报价汇总", lambda: parse_haitou(wb["海派头程报价汇总"])),
        ("集运价格汇总", lambda: parse_jiyun(wb["集运价格汇总"])),
        ("HYE海派渠道汇总", lambda: parse_haipai_channels(wb["HYE海派渠道汇总"])),
        ("HYE快船海卡", lambda: parse_haika_sheet(wb["HYE快船海卡"], "HYE快船海卡", "HYE快船海卡")),
        ("HYE普船海卡", lambda: parse_haika_sheet(wb["HYE普船海卡"], "HYE普船海卡", "HYE普船海卡")),
        ("HYE美西特惠普船LA海卡", lambda: parse_haika_sheet(wb["HYE美西特惠普船LA海卡"], "HYE美西特惠普船LA海卡", "HYE美西特惠普船LA海卡")),
        ("价格汇总", lambda: parse_price_summary(wb["价格汇总"])),
        ("HYE美西商私卡", lambda: parse_shangsika(wb["HYE美西商私卡"], "HYE美西商私卡")),
        ("HYE美西商私卡自税汇总", lambda: parse_shangsika_zishui(wb["HYE美西商私卡自税汇总"])),
        ("HYE美中美东商私卡", lambda: parse_shangsika(wb["HYE美中美东商私卡"], "HYE美中美东商私卡")),
        ("HYE超大件", lambda: parse_zone_simple(wb["HYE超大件"], "HYE超大件", "超大件")),
        ("HYE小包裹", lambda: parse_zone_simple(wb["HYE小包裹"], "HYE小包裹", "小包裹")),
        ("HYE岛屿海派", lambda: parse_zone_simple(wb["HYE岛屿海派"], "HYE岛屿海派", "岛屿海派")),
        # Air book (2026-10)
        ("HYE空派", lambda: parse_kongpai(wb["HYE空派"], "HYE空派", "空派", "包税")),
        ("HYE空派自税", lambda: parse_kongpai(wb["HYE空派自税"], "HYE空派自税", "空派自税", "自税")),
        ("HYE空卡", lambda: parse_kongpai(wb["HYE空卡"], "HYE空卡", "空卡", "包税")),
        ("HYE空派小包", lambda: parse_kongpai_xiaobao(wb["HYE空派小包"])),
        ("HYE空派超大件", lambda: parse_kongpai(wb["HYE空派超大件"], "HYE空派超大件", "空派超大件", "包税")),
    ]

    sheet_stats = {}
    for name, fn in parsers:
        if name not in wb.sheetnames:
            sheet_stats[name] = "SKIP: sheet missing"
            print(f"  {name}: SKIP (sheet missing)")
            continue
        try:
            rows = fn()
            sheet_stats[name] = len(rows)
            all_rates.extend(rows)
            print(f"  {name}: {len(rows)} rates")
        except Exception as e:
            sheet_stats[name] = f"ERROR: {e}"
            print(f"  {name}: ERROR {e}")

    schedule = []
    if "船期表" in wb.sheetnames:
        schedule = parse_schedule(wb["船期表"])
        print(f"  船期表: {len(schedule)} rows")
    else:
        print("  船期表: SKIP (sheet missing)")

    surcharges = []
    if "海运品名附加费" in wb.sheetnames:
        surcharges = parse_surcharges(wb["海运品名附加费"])
        print(f"  海运品名附加费: {len(surcharges)} items")
    if "美国海外仓快速中转服务" in wb.sheetnames:
        extra = parse_warehouse_surcharges(wb["美国海外仓快速中转服务"])
        surcharges.extend(extra)
        print(f"  美国海外仓快速中转服务: {len(extra)} items")

    all_rates = assign_ids(all_rates)
    meta = build_meta(catalog, all_rates, schedule, surcharges)
    meta["company"] = "AQ国际"
    meta["brand"] = "AQ"
    meta["sheet_stats"] = sheet_stats

    # Drop stale per-category files from previous books
    if OUT.exists():
        for old in OUT.glob("rates_*.json"):
            old.unlink()

    save_json("rates.json", all_rates)
    save_json("meta.json", meta)
    save_json("index.json", {
        "company": "AQ国际",
        "brand": "AQ",
        "categories": meta["categories"],
        "products_by_category": meta["products_by_category"],
        "warehouses": meta["warehouses"],
        "tax_modes": meta["tax_modes"],
        "fba_count": len(meta["fba_warehouse_codes"]),
        "rate_count": meta["rate_count"],
        "disclaimer": meta["disclaimer"],
        "parsed_at": meta["parsed_at"],
        "zip_zones": meta["zip_zones"],
    })
    from collections import defaultdict
    by_cat = defaultdict(list)
    for r in all_rates:
        by_cat[r.get("category") or "其他"].append(r)
    for cat, rows in by_cat.items():
        save_json(f"rates_{cat}.json", rows)
    fba_codes = sorted({
        c.upper()
        for r in all_rates
        for c in (r.get("warehouse_codes") or [])
        if c and "海外仓" not in c and len(c) <= 16
    })
    save_json("fba_codes.json", fba_codes)
    print(f"\nDone. Total rates: {len(all_rates)}")



if __name__ == "__main__":
    main()
